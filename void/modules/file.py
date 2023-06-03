from .color_code import bcolors as bc

from openpyxl import load_workbook
import pandas as pd


class File():

    def __init__(self, file_name, sheet_name, load_dir):
        self.file_name = file_name
        self.sheet_name = sheet_name
        self.load_dir = load_dir

    def import_excel(self):
        table_list = []
        workbook = load_workbook(self.load_dir, data_only='True')

        worksheet = workbook[self.sheet_name]

        mapping = {}

        # loop through all the tables and add to a dictionary
        for table, data_boundary in worksheet.tables.items():
            # parse the data within the ref boundary
            data = worksheet[data_boundary]

            ### extract the data ###
            # the inner list comprehension gets the values for each cell in the table
            content = [[cell.value for cell in ent]
                       for ent in data]
            header = content[0]

            # the contents ... excluding the header
            rest = content[1:]

            # create dataframe with the column names
            # and pair table name with dataframe
            df = pd.DataFrame(rest, columns=header)
            mapping[table] = df
            table_list.append(mapping[table])
            self.file_name = self.file_name.replace('_', ' ')
            print(
                f'{bc.HEADER}{bc.BOLD}{self.file_name.capitalize()}{bc.ENDC} ' f'{bc.OKGREEN}{table} Loaded{bc.ENDC}')
            
        
    def export_excel(self):
        
        return table_list

